"""Flatten the inventory into one wide table — every published field is a column.

Default: *all* columns. A fixed set of "core" columns comes first (the
normalised view HERO computes across sources: source, number, type, year,
status and where that status came from, ISO dates, links, KB location),
followed by every raw field any source publishes, labelled exactly as the
source prints it. A record from a source that lacks a field shows it blank.

Formats: CSV (UTF-8 with BOM so Excel reads Indonesian characters),
XLSX, JSON, and a single self-contained HTML page for browsing.
"""
from __future__ import annotations

import csv
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

# Control characters (other than tab/newline) that Excel's XML format cannot
# encode at all — OJK's source pages occasionally embed these in scraped text.
_ILLEGAL_XLSX_CHARS = re.compile(
    "[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")

from hero.config import Settings
from hero.kb.catalog import Catalog

SOURCE_ORDER = ["jdih-ojk", "ojk-regulasi", "ojk-rancangan"]
SOURCE_LABELS = {
    "jdih-ojk": "JDIH OJK",
    "ojk-regulasi": "OJK · Regulasi",
    "ojk-rancangan": "OJK · Rancangan Regulasi",
}
SOURCE_URLS = {
    "jdih-ojk": "https://jdih.ojk.go.id/Web/ViewPeraturan/Index",
    "ojk-regulasi": "https://ojk.go.id/id/regulasi/default.aspx",
    "ojk-rancangan": "https://www.ojk.go.id/id/regulasi/otoritas-jasa-keuangan/"
                     "rancangan-regulasi/Default.aspx",
}
STATUS_LABELS = {
    "berlaku": "Berlaku", "diubah": "Diubah", "dicabut": "Dicabut",
    "rancangan": "Rancangan", "unknown": "Tidak diketahui",
}

# (key, label) — computed per record in _core_values().
CORE_COLUMNS: list[tuple[str, str]] = [
    ("source", "Sumber"),
    ("title", "Judul"),
    ("number", "Nomor"),
    ("doc_type", "Jenis (singkatan)"),
    ("jenis", "Jenis"),
    ("sektor", "Sektor"),
    ("category", "Kategori HERO"),
    ("year", "Tahun"),
    ("status", "Status"),
    ("status_label", "Status (tulisan sumber)"),
    ("status_source", "Status berasal dari"),
    ("date_penetapan", "Tanggal Penetapan"),
    ("date_pengundangan", "Tanggal Pengundangan"),
    ("date_berlaku", "Tanggal Berlaku"),
    ("date_rancangan", "Tanggal Rancangan Dipublikasikan"),
    ("document_name", "Nama Dokumen"),
    ("doc_format", "Format Dokumen"),
    ("attachment_count", "Jumlah Berkas Terkait"),
    ("attachments", "Berkas Terkait"),
    ("in_kb", "Di Knowledge Base"),
    ("kb_path", "Lokasi di Knowledge Base"),
    ("detail_url", "Tautan Halaman"),
    ("document_url", "Tautan Dokumen"),
    ("listed_at", "Pertama Terdaftar"),
    ("enriched_at", "Detail Dibaca"),
]
N_CORE = len(CORE_COLUMNS)

# Raw fields HERO derived itself and already exposes as a core column.
_DERIVED_RAW = {"Tanggal Penetapan (ISO)", "Tanggal Pengundangan (ISO)",
                "Tanggal Berlaku (ISO)", "Tanggal (ISO)"}
_STATUS_SOURCE_LABELS = {"jdih": "JDIH OJK", "sumber": "situs sumber",
                         "teks-dokumen": "teks dokumen"}
_EXCEL_CELL_LIMIT = 32_000


def _doc_format(row: dict[str, Any], attachments: list[dict]) -> str | None:
    """What HERO can do with this record's document, in one word or two."""
    name = (row.get("document_name") or "").lower()
    url = (row.get("document_url") or "").lower().split("?")[0]
    if url:
        if name.endswith(".zip") or url.endswith(".zip"):
            return "zip → pdf"
        return "pdf"
    exts = {a.get("ext") for a in attachments if a.get("ext")}
    if exts & {"docx", "doc"}:
        return "docx — di luar cakupan PDF"
    return None


def _core_values(row: dict[str, Any], fields: dict[str, Any],
                 attachments: list[dict]) -> dict[str, Any]:
    atts = ["{}: {}".format(a.get("kind", "berkas"), a.get("name", "")) for a in attachments]
    return {
        "source": SOURCE_LABELS.get(row["source"], row["source"]),
        "title": row["title"],
        "number": row["number"],
        "doc_type": row["doc_type"],
        "jenis": row["jenis"],
        "sektor": row["sektor"],
        "category": row["category"],
        "year": row["year"],
        "status": STATUS_LABELS.get(row["status"] or "unknown", row["status"]),
        "status_label": row["status_label"],
        "status_source": _STATUS_SOURCE_LABELS.get(row["status_source"] or "",
                                                   row["status_source"]),
        "date_penetapan": fields.get("Tanggal Penetapan (ISO)"),
        "date_pengundangan": fields.get("Tanggal Pengundangan (ISO)"),
        "date_berlaku": fields.get("Tanggal Berlaku (ISO)"),
        "date_rancangan": fields.get("Tanggal (ISO)"),
        "document_name": row["document_name"],
        "doc_format": _doc_format(row, attachments),
        "attachment_count": len(attachments) or None,
        "attachments": "; ".join(atts) or None,
        "in_kb": "ya" if row["doc_id"] else "belum",
        "kb_path": row["kb_path"],
        "detail_url": row["detail_url"],
        "document_url": row["document_url"],
        "listed_at": row["listed_at"],
        "enriched_at": row["enriched_at"],
    }


def build_table(cat: Catalog, columns: Sequence[str] | None = None,
                sources: Sequence[str] | None = None) -> tuple[list[str], list[list[Any]]]:
    """Return (header labels, rows). ``columns`` filters by label (case-insensitive)."""
    params: list[Any] = []
    where = ""
    if sources:
        where = f"WHERE i.source IN ({', '.join('?' for _ in sources)})"
        params.extend(sources)
    raw_rows = cat.conn.execute(
        f"""SELECT i.*, d.stored_path AS kb_path FROM inventory i
            LEFT JOIN documents d ON d.doc_id = i.doc_id {where}""", params).fetchall()
    order = {s: n for n, s in enumerate(SOURCE_ORDER)}
    raw_rows = sorted(raw_rows, key=lambda r: (order.get(r["source"], 99),
                                               -(r["year"] or 0), r["title"] or ""))

    # Raw field labels in first-seen order, so each source's own sequence
    # of fields stays together in the export.
    parsed: list[tuple[Any, dict, list]] = []
    raw_labels: list[str] = []
    seen: set[str] = set()
    core_labels = {label.lower() for _, label in CORE_COLUMNS}
    for r in raw_rows:
        fields = json.loads(r["fields_json"] or "{}")
        atts = json.loads(r["attachments_json"] or "[]")
        parsed.append((r, fields, atts))
        for key in fields:
            if key not in seen and key not in _DERIVED_RAW:
                seen.add(key)
                raw_labels.append(key)

    def raw_header(label: str) -> str:
        # A raw field whose label matches a core column keeps its data under
        # a distinct header rather than being dropped.
        return f"{label} (sumber)" if label.lower() in core_labels else label

    header = [label for _, label in CORE_COLUMNS] + [raw_header(k) for k in raw_labels]
    rows: list[list[Any]] = []
    for r, fields, atts in parsed:
        core = _core_values(dict(r), fields, atts)
        rows.append([core[key] for key, _ in CORE_COLUMNS]
                    + [fields.get(k) for k in raw_labels])

    if columns:
        wanted = {c.strip().lower() for c in columns}
        keep = [i for i, h in enumerate(header) if h.lower() in wanted]
        header = [header[i] for i in keep]
        rows = [[row[i] for i in keep] for row in rows]
    return header, rows


def write_csv(path: Path, header: list[str], rows: list[list[Any]]) -> Path:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows([["" if v is None else v for v in row] for row in rows])
    return path


def write_json(path: Path, header: list[str], rows: list[list[Any]]) -> Path:
    records = [{h: v for h, v in zip(header, row) if v not in (None, "")} for row in rows]
    path.write_text(json.dumps(records, ensure_ascii=False, indent=1), encoding="utf-8")
    return path


def write_xlsx(path: Path, header: list[str], rows: list[list[Any]]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventaris"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F4FA8")
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    def clean(v: Any) -> Any:
        if not isinstance(v, str):
            return v
        return _ILLEGAL_XLSX_CHARS.sub(" ", v)[:_EXCEL_CELL_LIMIT]

    for row in rows:
        ws.append([clean(v) for v in row])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, start=1):
        sample = [len(str(r[i - 1])) for r in rows[:300] if r[i - 1] is not None]
        width = min(max([len(h)] + sample) + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = max(width, 10)
    wb.save(path)
    return path


def _summary(cat: Catalog) -> dict[str, Any]:
    c = cat.conn
    by_source = {r[0]: r[1] for r in c.execute(
        "SELECT source, COUNT(*) FROM inventory GROUP BY source")}
    by_status = {r[0] or "unknown": r[1] for r in c.execute(
        "SELECT status, COUNT(*) FROM inventory GROUP BY status")}
    enriched = {r[0]: r[1] for r in c.execute(
        "SELECT source, COUNT(*) FROM inventory WHERE enriched_at IS NOT NULL GROUP BY source")}
    in_kb = {r[0]: r[1] for r in c.execute(
        "SELECT source, COUNT(*) FROM inventory WHERE doc_id IS NOT NULL GROUP BY source")}
    return {
        "sources": [
            {"key": s, "label": SOURCE_LABELS.get(s, s), "url": SOURCE_URLS.get(s),
             "count": by_source.get(s, 0), "enriched": enriched.get(s, 0),
             "in_kb": in_kb.get(s, 0)}
            for s in SOURCE_ORDER if s in by_source
        ] + [
            {"key": s, "label": s, "url": None, "count": n,
             "enriched": enriched.get(s, 0), "in_kb": in_kb.get(s, 0)}
            for s, n in by_source.items() if s not in SOURCE_ORDER
        ],
        "status": [{"key": k, "label": STATUS_LABELS.get(k, k), "count": by_status.get(k, 0)}
                   for k in ("berlaku", "diubah", "dicabut", "rancangan", "unknown")
                   if by_status.get(k)],
        "total": sum(by_source.values()),
    }


def write_html(path: Path, header: list[str], rows: list[list[Any]],
               summary: dict[str, Any], title: str = "Register Peraturan HERO") -> Path:
    template = (Path(__file__).parent / "templates" / "inventory_page.html").read_text(
        encoding="utf-8")
    payload = {
        "columns": header, "core": N_CORE, "rows": rows, "summary": summary,
        "generated": datetime.now().isoformat(timespec="minutes"),
    }
    data = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    # Neutralise "</script>" inside the data so it cannot close the tag early.
    data = data.replace("</", "<\\/")
    html = template.replace("__TITLE__", title).replace("__DATA__", data)
    path.write_text(html, encoding="utf-8")
    return path


def export_inventory(
    settings: Settings, out_dir: Path,
    formats: Sequence[str] = ("csv", "xlsx", "json", "html"),
    columns: Sequence[str] | None = None, sources: Sequence[str] | None = None,
) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    with Catalog(settings.catalog_db) as cat:
        header, rows = build_table(cat, columns, sources)
        summary = _summary(cat)
    written: dict[str, Path] = {}
    base = out_dir / "inventaris_peraturan"
    for fmt in formats:
        fmt = fmt.lower()
        if fmt == "csv":
            written["csv"] = write_csv(base.with_suffix(".csv"), header, rows)
        elif fmt == "json":
            written["json"] = write_json(base.with_suffix(".json"), header, rows)
        elif fmt == "xlsx":
            written["xlsx"] = write_xlsx(base.with_suffix(".xlsx"), header, rows)
        elif fmt == "html":
            written["html"] = write_html(base.with_suffix(".html"), header, rows, summary)
        else:
            raise ValueError(f"unknown export format: {fmt}")
    return written
