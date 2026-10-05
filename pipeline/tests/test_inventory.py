"""Inventory: normalisation, record building, merge rules, status reconciliation,
export, archive handling and KB restructuring."""
from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from bs4 import BeautifulSoup

from hero import inventory as inv
from hero.ingest import jdih
from hero.ingest import sharepoint as sp
from hero.ingest.archive import extract_primary_pdf, is_zip
from hero.kb.catalog import Catalog

FIX = Path(__file__).parent / "fixtures" / "html"


# -- normalisation ---------------------------------------------------------
@pytest.mark.parametrize("text,number", [
    ("Peraturan OJK Nomor 9/POJK.04/2015 tentang Repo", "9/POJK.04/2015"),
    ("Undang-Undang Nomor 8 Tahun 1995 tentang Pasar Modal", "8 Tahun 1995"),
    ("SEOJK Nomor 20 /SEOJK.08/ 2025 tentang X", "20/SEOJK.08/2025"),
    ("tanpa nomor", None),
])
def test_number_from_text(text, number):
    assert inv.number_from_text(text) == number


@pytest.mark.parametrize("number,abbrev", [
    ("9/POJK.04/2015", "POJK"), ("45/PADK.06/2025", "PADK"),
    ("2/11/PBI/2000", "PBI"), ("8 Tahun 1995", None), (None, None),
])
def test_abbrev_from_number(number, abbrev):
    assert inv.abbrev_from_number(number) == abbrev


def test_reg_key_agrees_across_number_styles():
    """The same regulation written two ways must produce one key."""
    a = inv.reg_key("POJK", "9/POJK.04/2015", 2015)
    b = inv.reg_key("pojk", "9", 2015)
    assert a == b == "POJK|9|2015"
    assert inv.reg_key(None, "9", 2015) is None
    assert inv.reg_key("POJK", "9", None) is None


@pytest.mark.parametrize("raw,iso", [("25-06-2015", "2015-06-25"), ("1-1-2020", "2020-01-01"),
                                     ("2015-06-25", None), ("", None)])
def test_dmy_to_iso(raw, iso):
    assert inv.dmy_to_iso(raw) == iso


def test_draft_type_from_url_or_text():
    assert inv.draft_type("https://x/Pages/RPADK-Laporan.aspx") == "RPADK"
    assert inv.draft_type("https://x/Pages/Laporan.aspx",
                          "Rancangan Surat Edaran OJK tentang X") == "RSEOJK"
    assert inv.draft_type("https://x/Pages/Laporan.aspx", "Laporan") is None


def test_source_key_for_ref():
    assert inv.source_key_for_ref("web", "https://jdih.ojk.go.id/Web/x") == "jdih-ojk"
    assert inv.source_key_for_ref(
        "web", "https://www.ojk.go.id/id/regulasi/otoritas-jasa-keuangan/"
               "rancangan-regulasi/Documents/a.zip") == "ojk-rancangan"
    assert inv.source_key_for_ref("web", "https://ojk.go.id/id/regulasi/Documents/a.pdf") \
        == "ojk-regulasi"
    assert inv.source_key_for_ref("upload", "/tmp/a.pdf") == "unggah-manual"
    assert inv.source_key_for_ref("web", "https://www.bi.go.id/a.pdf") == "bi-go-id"


# -- JDIH detail: several labelled documents -------------------------------
MULTI_DOC = """
<table>
 <tr><th>Judul</th><td>:</td><td>PADK Nomor 1 Tahun 2026</td></tr>
 <tr><th>Singkatan Jenis/Bentuk Peraturan</th><td>:</td><td>PADK</td></tr>
 <tr><th>Tanggal Penetapan</th><td>:</td><td>05-01-2026</td></tr>
 <tr><th>Status Peraturan</th><td>:</td><td>Berlaku</td></tr>
</table>
<table>
 <tr><td>Abstrak</td><td>:</td><td><a href="/Web/ViewPeraturan/PreviewDokumen/a1">2026abspadk001.pdf</a>
   <a href="/Web/ViewPeraturan/DownloadDokumen/a1">Unduh</a> * Klik pada nama file</td></tr>
 <tr><td>Peraturan</td><td>:</td><td><a href="/Web/ViewPeraturan/PreviewDokumen/p1">2026padk001.pdf</a>
   <a href="/Web/ViewPeraturan/DownloadDokumen/p1">Unduh</a> * Klik pada nama file</td></tr>
 <tr><td>FAQ</td><td>:</td><td><a href="/Web/ViewPeraturan/PreviewDokumen/f1">2026faqpadk001.pdf</a>
   <a href="/Web/ViewPeraturan/DownloadDokumen/f1">Unduh</a></td></tr>
</table>
"""


def test_jdih_detail_primary_is_the_regulation_not_the_abstract():
    d = jdih.parse_detail(MULTI_DOC)
    assert d["document_url"].endswith("/DownloadDokumen/p1")
    assert d["document_name"] == "2026padk001.pdf"
    assert len(d["documents"]) == 3
    # Document rows hold the file name, not the page's instructions.
    assert d["fields"]["Abstrak"] == "2026abspadk001.pdf"
    assert "Klik" not in d["fields"]["Peraturan"]


def test_jdih_detail_real_page():
    d = jdih.parse_detail((FIX / "jdih_detail.html").read_text(encoding="utf-8"))
    assert d["fields"]["Tanggal Penetapan"] == "25-06-2015"
    assert d["fields"]["Status Peraturan"] == "Berlaku"
    assert len(d["landasan"]) == 4
    assert d["riwayat"].startswith("Mencabut")
    assert d["document_name"] == "pojk 9-2015.pdf"


def test_apply_jdih_detail_builds_attachments_with_kinds():
    entry = jdih.JdihEntry(detail_url="https://jdih.ojk.go.id/d/1/01/09",
                           title="PADK Nomor 1 Tahun 2026 tentang X", sektor="01",
                           jenis="09", status="Berlaku")
    rec = inv.record_from_jdih(entry)
    upd = inv.apply_jdih_detail(rec, jdih.parse_detail(MULTI_DOC))
    kinds = {a["name"]: a["kind"] for a in upd["attachments"]}
    assert kinds["2026padk001.pdf"] == "utama"
    assert kinds["2026abspadk001.pdf"] == "abstrak"
    assert kinds["2026faqpadk001.pdf"] == "faq"
    assert upd["doc_type"] == "PADK"
    assert upd["fields"]["Tanggal Penetapan (ISO)"] == "2026-01-05"
    assert upd["reg_key"] == "PADK|1|2026"


# -- SharePoint rows → records ----------------------------------------------
def test_regulasi_record_from_real_page():
    s = BeautifulSoup((FIX / "ojk_regulasi_list.html").read_text(encoding="utf-8"), "lxml")
    row = sp.parse_regulasi_rows(s, "https://ojk.go.id/id/regulasi/default.aspx", 1)[0]
    rec = inv.record_from_sharepoint_row(row, inv.SOURCE_REGULASI)
    assert rec["doc_type"] == "PADK" and rec["year"] == 2025   # issue year, not the caption's 2027
    assert rec["category"] == "iknb"
    assert rec["status"] == "unknown"                           # ojk.go.id never says


def test_rancangan_record_is_always_rancangan():
    s = BeautifulSoup((FIX / "ojk_rancangan_list.html").read_text(encoding="utf-8"), "lxml")
    row = sp.parse_rancangan_rows(s, "https://www.ojk.go.id/x/rancangan-regulasi/Default.aspx", 1)[0]
    rec = inv.record_from_sharepoint_row(row, inv.SOURCE_RANCANGAN)
    assert rec["status"] == "rancangan"
    assert rec["doc_type"] == "RPADK"


# -- catalog merge rules ------------------------------------------------------
def test_relisting_never_erases_detail_or_document_link(tmp_path):
    cat = Catalog(tmp_path / "c.db")
    entry = jdih.JdihEntry(detail_url="https://jdih.ojk.go.id/d/1/01/06",
                           title="POJK Nomor 5/POJK.03/2020 tentang X", sektor="01",
                           jenis="06", status="Berlaku")
    assert cat.upsert_inventory(inv.record_from_jdih(entry)) is True
    cat.upsert_inventory({"record_key": entry.detail_url, "fields": {"Subjek": "X"},
                          "document_url": "https://jdih.ojk.go.id/dl/1",
                          "enriched_at": "2026-09-15T10:00:00"})
    cat.link_inventory_document(entry.detail_url, "doc-123")
    assert cat.upsert_inventory(inv.record_from_jdih(entry)) is False   # listing again
    row = cat.get_inventory(entry.detail_url)
    assert json.loads(row["fields_json"])["Subjek"] == "X"
    assert row["document_url"] == "https://jdih.ojk.go.id/dl/1"
    assert row["doc_id"] == "doc-123"
    cat.close()


def test_reconcile_copies_jdih_status_onto_ojk_rows(tmp_settings):
    with Catalog(tmp_settings.catalog_db) as cat:
        cat.upsert_inventory({
            "record_key": "jdih/1", "source": inv.SOURCE_JDIH, "status": "dicabut",
            "status_label": "Tidak Berlaku", "status_source": "jdih",
            "reg_key": "POJK|12|2016", "detail_url": "https://jdih/1"})
        cat.upsert_inventory({
            "record_key": "ojk/1", "source": inv.SOURCE_REGULASI, "status": "unknown",
            "reg_key": "POJK|12|2016"})
        cat.upsert_inventory({
            "record_key": "ojk/2", "source": inv.SOURCE_REGULASI, "status": "unknown",
            "reg_key": "POJK|99|2016"})
    result = inv.reconcile_status(tmp_settings)
    assert result == {"matched": 1, "unmatched": 1}
    with Catalog(tmp_settings.catalog_db) as cat:
        row = cat.get_inventory("ojk/1")
        assert row["status"] == "dicabut" and row["status_source"] == "jdih"
        assert cat.get_inventory("ojk/2")["status"] == "unknown"


# -- export ---------------------------------------------------------------------
def test_export_has_every_field_as_a_column(tmp_settings, tmp_path):
    from hero.kb.inventory_export import N_CORE, build_table, export_inventory

    with Catalog(tmp_settings.catalog_db) as cat:
        cat.upsert_inventory({"record_key": "a", "source": inv.SOURCE_JDIH, "title": "A",
                              "status": "berlaku",
                              "fields": {"Bidang Hukum": "Ekonomi", "Judul": "A (asli)",
                                         "Tanggal Penetapan (ISO)": "2015-06-25"}})
        cat.upsert_inventory({"record_key": "b", "source": inv.SOURCE_REGULASI, "title": "B",
                              "status": "unknown", "fields": {"SubSektor": "Bank Umum"}})
        header, rows = build_table(cat)
    assert "Bidang Hukum" in header and "SubSektor" in header
    assert "Judul (sumber)" in header          # collision with a core label is kept, renamed
    assert "Tanggal Penetapan (ISO)" not in header   # already a core column
    assert header.index("Bidang Hukum") >= N_CORE
    row_a = rows[[r[1] for r in rows].index("A")]
    assert row_a[header.index("Tanggal Penetapan")] == "2015-06-25"

    written = export_inventory(tmp_settings, tmp_path / "out")
    assert set(written) == {"csv", "xlsx", "json", "html"}
    page = written["html"].read_text(encoding="utf-8")
    assert "__DATA__" not in page and "Bidang Hukum" in page


def test_export_column_selection(tmp_settings):
    from hero.kb.inventory_export import build_table

    with Catalog(tmp_settings.catalog_db) as cat:
        cat.upsert_inventory({"record_key": "a", "source": inv.SOURCE_JDIH, "title": "A"})
        header, rows = build_table(cat, columns=["judul", "Status"])
    assert header == ["Judul", "Status"]
    assert rows[0][0] == "A"


# -- archives --------------------------------------------------------------------
def test_extract_primary_pdf_prefers_batang_tubuh(tmp_path):
    z = tmp_path / "draft.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("RPADK/02 Lampiran Besar.pdf", b"%PDF-1.4 " + b"x" * 5000)
        zf.writestr("RPADK/00 Batang Tubuh.pdf", b"%PDF-1.4 body")
        zf.writestr("RPADK/matriks tanggapan.docx", b"docx")
    assert is_zip(z)
    pdf, members = extract_primary_pdf(z, tmp_path / "out")
    assert pdf is not None and "Batang Tubuh" in pdf.name
    assert pdf.read_bytes().startswith(b"%PDF")
    assert len(members) == 3


def test_extract_primary_pdf_is_zip_slip_safe(tmp_path):
    z = tmp_path / "evil.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("../../escape.pdf", b"%PDF-1.4 x")
    pdf, _ = extract_primary_pdf(z, tmp_path / "out")
    assert pdf is not None
    assert pdf.parent == tmp_path / "out"
    assert not (tmp_path.parent / "escape.pdf").exists()


def test_extract_primary_pdf_without_pdf(tmp_path):
    z = tmp_path / "docs.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("a.docx", b"docx")
    pdf, members = extract_primary_pdf(z, tmp_path / "out")
    assert pdf is None and len(members) == 1


def test_docx_only_record_is_labelled_and_not_harvested(tmp_settings):
    from hero.kb.inventory_export import build_table

    with Catalog(tmp_settings.catalog_db) as cat:
        cat.upsert_inventory({
            "record_key": "draft/1", "source": inv.SOURCE_RANCANGAN, "title": "RPOJK BPRS",
            "status": "rancangan", "enriched_at": "2026-09-15T10:00:00",
            "attachments": [{"name": "RPOJK BPRS.docx", "url": "https://x/a.docx",
                             "kind": "utama", "ext": "docx"}]})
        cat.upsert_inventory({
            "record_key": "reg/1", "source": inv.SOURCE_REGULASI, "title": "POJK 1",
            "enriched_at": "2026-09-15T10:00:00", "document_url": "https://x/p.pdf",
            "document_name": "p.pdf"})
        pending = [r["record_key"] for r in cat.list_inventory(pending_download=True)]
        assert pending == ["reg/1"]                 # the .docx-only draft is skipped
        header, rows = build_table(cat)
    fmt = {r[header.index("Judul")]: r[header.index("Format Dokumen")] for r in rows}
    assert fmt["RPOJK BPRS"].startswith("docx")
    assert fmt["POJK 1"] == "pdf"


def test_xlsx_export_strips_illegal_control_characters(tmp_settings, tmp_path):
    from hero.kb.inventory_export import write_xlsx

    header = ["Judul"]
    rows = [["Teks dengan \x0b karakter kontrol ilegal \x00 di tengah"]]
    path = write_xlsx(tmp_path / "t.xlsx", header, rows)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    val = wb.active.cell(row=2, column=1).value
    assert "\x0b" not in val and "\x00" not in val
    assert "karakter kontrol ilegal" in val
