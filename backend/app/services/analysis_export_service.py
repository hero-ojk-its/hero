"""
app/services/analysis_export_service.py
Layanan untuk mengekspor hasil analisa dokumen ke format DOCX, CSV, dan JSON (Langkah 22, US-40 #62).
"""
import csv
from datetime import datetime, timezone
import io
import json
import logging
import re
from typing import Any, Dict, Optional, Set, Tuple
import unicodedata
import urllib.parse

import docx
from docx.shared import Inches, Pt, RGBColor
from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.models.analysis_result import AnalysisResult
from app.models.analysis_run import AnalysisRun
from app.models.document import Document
from app.models.enums import JenisAnalisa, ModeAnalisa
from app.models.user import User
from app.services.analysis_result_service import get_document_fingerprint

logger = logging.getLogger("hero")

LABEL_REKOMENDASI = "DRAFT / REKOMENDASI — hasil analisa otomatis, wajib diverifikasi analis."


def format_pasal_ref(pasal: Any) -> str:
    """
    Format rujukan pasal untuk Key Takeaways.
    Bila nilai berupa teks lengkap (mis. 'Pasal 4 ayat (1)' atau 'Bagian II'), dipakai apa adanya.
    Awalan 'Pasal ' hanya ditambahkan bila nilainya murni nomor/huruf (mis. '9', '4A') untuk data lama.
    """
    if pasal is None:
        return ""
    p_str = str(pasal).strip()
    if not p_str:
        return ""
    if re.match(r"^\d+[A-Z]?$", p_str):
        return f"Pasal {p_str}"
    return p_str


def generate_export_filename(
    identitas: Optional[Dict[str, Any]],
    document_id: int,
    mode: str,
    ext: str,
) -> Tuple[str, str]:
    """
    Format: analisa_{jenis}_{nomor}_{tahun}_{mode}.{ext}, dari identitas.
    Bila kosong, pakai analisa_dokumen_{document_id}_{mode}.{ext}.
    Mengembalikan tuple (filename_ascii, filename_utf8).
    filename_ascii dibersihkan dari non-ASCII dan karakter tidak aman untuk header HTTP.
    filename_utf8 mempertahankan karakter Unicode asli untuk filename* (RFC 5987).
    """
    jenis = None
    nomor = None
    tahun = None
    if isinstance(identitas, dict):
        jenis = identitas.get("jenis")
        nomor = identitas.get("nomor")
        tahun = identitas.get("tahun")

    if jenis and nomor and tahun:
        raw_name = f"analisa_{jenis}_{nomor}_{tahun}_{mode}.{ext}"
    else:
        raw_name = f"analisa_dokumen_{document_id}_{mode}.{ext}"

    # UTF-8: bersihkan karakter yang tidak aman untuk nama berkas
    cleaned_utf8 = re.sub(r'[\s/\\:*?"<>|]+', "_", raw_name)
    cleaned_utf8 = re.sub(r"_+", "_", cleaned_utf8)

    # ASCII: normalisasi NFKD lalu buang karakter non-ASCII, kemudian bersihkan
    norm_ascii = unicodedata.normalize("NFKD", raw_name).encode("ascii", "ignore").decode("ascii")
    cleaned_ascii = re.sub(r'[\s/\\:*?"<>|]+', "_", norm_ascii)
    cleaned_ascii = re.sub(r"_+", "_", cleaned_ascii)
    if not cleaned_ascii or cleaned_ascii == f".{ext}":
        cleaned_ascii = f"analisa_dokumen_{document_id}_{mode}.{ext}"

    return cleaned_ascii, cleaned_utf8


def export_to_docx(
    doc_model: Document,
    result: Dict[str, Any],
    mode: str,
    ml_version: Optional[str],
    created_at: Optional[datetime],
    run_id: Optional[int],
    basi: bool,
    exported_by: Optional[str],
    dianalisa_oleh: Optional[str] = None,
) -> bytes:
    """Membangun dokumen DOCX hasil analisa sesuai spesifikasi US-40."""
    document = docx.Document()

    identitas = result.get("identitas") if isinstance(result.get("identitas"), dict) else {}
    kerangka = result.get("kerangka") if isinstance(result.get("kerangka"), list) else []
    ringkasan = result.get("ringkasan")
    ringkasan_teks = result.get("ringkasan_teks")
    ringkasan_ai = result.get("ringkasan_ai")
    poin_utama = result.get("poin_utama") if isinstance(result.get("poin_utama"), list) else []
    dasar_hukum = result.get("dasar_hukum") if isinstance(result.get("dasar_hukum"), list) else []
    sanksi = result.get("sanksi") if isinstance(result.get("sanksi"), list) else []
    catatan = result.get("catatan")

    # 1. Label besar di paling atas
    p_label = document.add_paragraph()
    run_label = p_label.add_run(LABEL_REKOMENDASI)
    run_label.bold = True
    run_label.font.size = Pt(13)
    run_label.font.color.rgb = RGBColor(180, 40, 40)

    # 2. Judul Dokumen
    jenis = identitas.get("jenis")
    nomor = identitas.get("nomor")
    tahun = identitas.get("tahun")
    tentang = identitas.get("tentang")

    parts = [str(x).strip() for x in [jenis, nomor, tahun] if x is not None and str(x).strip()]
    lead = " ".join(parts)
    if lead and tentang:
        judul_dokumen = f"{lead} tentang {tentang}"
    elif lead:
        judul_dokumen = lead
    elif tentang:
        judul_dokumen = str(tentang)
    else:
        judul_dokumen = doc_model.title or f"Dokumen Regulasi #{doc_model.id}"

    h_judul = document.add_heading(judul_dokumen, level=1)
    h_judul.style.font.size = Pt(16)

    # 3. Tabel Identitas
    document.add_heading("Identitas Regulasi", level=2)
    tbl_identitas = document.add_table(rows=0, cols=2)
    tbl_identitas.style = "Table Grid"

    fields_identitas = [
        ("Jenis", identitas.get("jenis")),
        ("Nomor", identitas.get("nomor")),
        ("Tahun", identitas.get("tahun")),
        ("Tentang", identitas.get("tentang")),
        ("Status", identitas.get("status")),
        ("Tanggal Terbit", identitas.get("tanggal_terbit")),
        ("Penerbit", identitas.get("penerbit")),
        ("Jumlah Pasal", identitas.get("jumlah_pasal")),
        ("Jumlah Bab", len(kerangka) if kerangka else "-"),
    ]

    for label, val in fields_identitas:
        row = tbl_identitas.add_row()
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].text = str(val) if (val is not None and str(val).strip()) else "-"

    document.add_paragraph()

    # 4. Tabel Informasi Eksekusi
    document.add_heading("Informasi Eksekusi Analisa", level=2)
    tbl_eksekusi = document.add_table(rows=0, cols=2)
    tbl_eksekusi.style = "Table Grid"

    waktu_dibuat_str = created_at.isoformat() if created_at else "-"
    waktu_ekspor_str = datetime.now(timezone.utc).isoformat()
    versi_aturan = ml_version or result.get("versi") or "-"

    fields_eksekusi = [
        ("Mode Dipakai", mode),
        ("Versi Aturan", versi_aturan),
        ("Dibuat", waktu_dibuat_str),
        ("Dianalisa Oleh", dianalisa_oleh or "-"),
        ("Run ID", str(run_id) if run_id is not None else "-"),
        ("Diekspor Pada", waktu_ekspor_str),
        ("Diekspor Oleh", exported_by or "-"),
    ]

    for label, val in fields_eksekusi:
        row = tbl_eksekusi.add_row()
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].bold = True
        row.cells[1].text = str(val)

    if basi:
        p_basi = document.add_paragraph()
        run_basi = p_basi.add_run("Dokumen telah berubah sejak analisa ini dibuat — jalankan analisa ulang.")
        run_basi.bold = True
        run_basi.font.color.rgb = RGBColor(200, 100, 0)

    document.add_paragraph()

    # 5. Ringkasan
    document.add_heading("Ringkasan", level=2)
    if isinstance(ringkasan, list) and ringkasan:
        for item in ringkasan:
            if isinstance(item, dict):
                kalimat = item.get("kalimat", "")
                rujukan = item.get("rujukan")
                if rujukan and isinstance(rujukan, list) and len(rujukan) > 0:
                    rujukan_str = ", ".join(str(r) for r in rujukan)
                    p_text = f"{kalimat} ({rujukan_str})"
                elif rujukan:
                    p_text = f"{kalimat} ({rujukan})"
                else:
                    p_text = kalimat
                document.add_paragraph(p_text)
            elif isinstance(item, str):
                document.add_paragraph(item)
    elif ringkasan_teks and isinstance(ringkasan_teks, str):
        document.add_paragraph(ringkasan_teks)
    else:
        document.add_paragraph("-")

    # Narasi AI-Assisted (setelah ringkasan deterministik)
    if mode == "ai" and ringkasan_ai and isinstance(ringkasan_ai, str):
        document.add_heading("Narasi AI-Assisted", level=3)
        document.add_paragraph(ringkasan_ai)

    # 6. Key Takeaways (Poin Utama)
    document.add_heading("Key Takeaways", level=2)
    if poin_utama:
        for idx, item in enumerate(poin_utama, start=1):
            if isinstance(item, dict):
                kategori = item.get("kategori") or "-"
                teks = item.get("teks") or "-"
                if item.get("terpotong"):
                    teks = f"{teks} (dipotong)"

                pasal = item.get("pasal")
                halaman = item.get("halaman")
                ref_parts = []
                if pasal:
                    pasal_formatted = format_pasal_ref(pasal)
                    if pasal_formatted:
                        ref_parts.append(pasal_formatted)
                if halaman is not None:
                    ref_parts.append(f"hal. {halaman}")

                ref_str = f" — {', '.join(ref_parts)}" if ref_parts else ""
                p_poin = document.add_paragraph()
                p_poin.add_run(f"{idx}. [{kategori}] {teks}{ref_str}")

                parafrase = item.get("parafrase")
                if parafrase and str(parafrase).strip():
                    p_par = document.add_paragraph()
                    p_par.paragraph_format.left_indent = Inches(0.25)
                    r_lbl = p_par.add_run("Parafrase AI: ")
                    r_lbl.bold = True
                    p_par.add_run(str(parafrase))
            elif isinstance(item, str):
                document.add_paragraph(f"{idx}. {item}")
    else:
        document.add_paragraph("-")

    # 7. Dasar Hukum dan Sanksi (lewati bila kosong)
    if dasar_hukum:
        document.add_heading("Dasar Hukum", level=2)
        for dh in dasar_hukum:
            document.add_paragraph(str(dh), style="List Bullet")

    if sanksi:
        document.add_heading("Sanksi", level=2)
        for s in sanksi:
            document.add_paragraph(str(s), style="List Bullet")

    # 8. Catatan Kaki & Tautan PDF
    document.add_heading("Catatan Tambahan & Tautan Dokumen", level=2)
    if catatan:
        document.add_paragraph(str(catatan))

    p_link = document.add_paragraph()
    p_link.add_run("Tautan Dokumen Asli (PDF): ").bold = True
    p_link.add_run(f"/api/v1/documents/{doc_model.id}/pdf")

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_to_csv(result: Dict[str, Any]) -> bytes:
    """Membangun berkas CSV (UTF-8 dengan BOM) hasil analisa sesuai spesifikasi US-40."""
    output = io.StringIO()
    writer = csv.writer(output)

    # Header
    writer.writerow(["bagian", "urutan", "kategori", "pasal_rujukan", "halaman", "teks", "parafrase"])

    # Baris pertama setelah header: label
    writer.writerow(["label", 1, "", "", "", LABEL_REKOMENDASI, ""])

    identitas = result.get("identitas") if isinstance(result.get("identitas"), dict) else {}
    kerangka = result.get("kerangka") if isinstance(result.get("kerangka"), list) else []
    ringkasan = result.get("ringkasan")
    poin_utama = result.get("poin_utama") if isinstance(result.get("poin_utama"), list) else []
    dasar_hukum = result.get("dasar_hukum") if isinstance(result.get("dasar_hukum"), list) else []
    sanksi = result.get("sanksi") if isinstance(result.get("sanksi"), list) else []

    # Bagian Identitas: satu baris per field
    idx_id = 1
    for k in ["jenis", "nomor", "tahun", "tentang", "status", "tanggal_terbit", "penerbit", "jumlah_pasal"]:
        val = identitas.get(k)
        writer.writerow(["identitas", idx_id, k, "", "", str(val) if val is not None else "", ""])
        idx_id += 1

    if kerangka:
        writer.writerow(["identitas", idx_id, "jumlah_bab", "", "", str(len(kerangka)), ""])

    # Bagian Ringkasan: satu baris per kalimat
    if isinstance(ringkasan, list) and ringkasan:
        for idx, item in enumerate(ringkasan, start=1):
            if isinstance(item, dict):
                kalimat = item.get("kalimat", "")
                rujukan = item.get("rujukan")
                if isinstance(rujukan, list):
                    pasal_rujukan = "; ".join(str(r) for r in rujukan)
                elif rujukan:
                    pasal_rujukan = str(rujukan)
                else:
                    pasal_rujukan = ""
                writer.writerow(["ringkasan", idx, "", pasal_rujukan, "", kalimat, ""])
            elif isinstance(item, str):
                writer.writerow(["ringkasan", idx, "", "", "", item, ""])
    elif result.get("ringkasan_teks"):
        writer.writerow(["ringkasan", 1, "", "", "", str(result.get("ringkasan_teks")), ""])

    # Bagian Poin Utama: satu baris per poin
    if isinstance(poin_utama, list) and poin_utama:
        for idx, item in enumerate(poin_utama, start=1):
            if isinstance(item, dict):
                kategori = item.get("kategori", "")
                pasal = item.get("pasal", "")
                halaman = item.get("halaman", "")
                teks = item.get("teks", "")
                if item.get("terpotong"):
                    teks = f"{teks} (dipotong)"
                parafrase = item.get("parafrase", "")
                writer.writerow([
                    "poin_utama",
                    idx,
                    kategori,
                    format_pasal_ref(pasal),
                    str(halaman) if halaman is not None else "",
                    teks,
                    str(parafrase) if parafrase is not None else "",
                ])
            elif isinstance(item, str):
                writer.writerow(["poin_utama", idx, "", "", "", item, ""])

    # Bagian Dasar Hukum: satu baris per butir
    if isinstance(dasar_hukum, list) and dasar_hukum:
        for idx, item in enumerate(dasar_hukum, start=1):
            writer.writerow(["dasar_hukum", idx, "", "", "", str(item), ""])

    # Bagian Sanksi: satu baris per butir
    if isinstance(sanksi, list) and sanksi:
        for idx, item in enumerate(sanksi, start=1):
            writer.writerow(["sanksi", idx, "", "", "", str(item), ""])

    # UTF-8 dengan BOM (agar terbaca langsung benar di Excel)
    csv_str = output.getvalue()
    return b"\xef\xbb\xbf" + csv_str.encode("utf-8")


def export_analysis(
    db: Session,
    doc: Document,
    format_type: str,
    mode: str,
    run_id: Optional[int],
    current_user: Optional[Any],
    allowed_ids: Set[int],
    saring_rujukan_fn: Any,
) -> Tuple[bytes, str, Dict[str, str], bool, str]:
    """
    Mengambil data analisa, menyaring rujukan, memformat ke docx/csv/json,
    dan mengembalikan (content_bytes, media_type, headers, basi, actual_mode).
    """
    raw_result = None
    actual_mode = mode
    ml_version = None
    created_at = None
    source_fingerprint = None
    created_by_user_id = None

    if run_id is not None:
        # Sumber 1: Dari riwayat eksekusi analysis_runs
        run = (
            db.query(AnalysisRun)
            .filter(AnalysisRun.id == run_id, AnalysisRun.document_id == doc.id)
            .first()
        )
        if not run:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Riwayat eksekusi analisa tidak ditemukan",
            )
        if run.sumber == "gagal":
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail={"code": "hasil_tidak_tersedia", "message": "Hasil analisa tidak tersedia untuk eksekusi ini."},
            )

        if run.result is not None:
            raw_result = run.result
        elif run.result_hash:
            # Resolusi dari run ML awal dengan result_hash sama
            src_run = (
                db.query(AnalysisRun)
                .filter(
                    AnalysisRun.document_id == doc.id,
                    AnalysisRun.result_hash == run.result_hash,
                    AnalysisRun.result.isnot(None),
                )
                .order_by(AnalysisRun.id.asc())
                .first()
            )
            if src_run:
                raw_result = src_run.result

        if raw_result is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail={"code": "hasil_tidak_tersedia", "message": "Hasil analisa tidak tersedia untuk eksekusi ini."},
            )

        actual_mode = (
            run.mode_dipakai.value
            if (run.mode_dipakai and hasattr(run.mode_dipakai, "value"))
            else (
                str(run.mode_dipakai)
                if run.mode_dipakai
                else (run.mode_diminta.value if hasattr(run.mode_diminta, "value") else str(run.mode_diminta))
            )
        )
        ml_version = run.ml_version
        created_at = run.created_at
        source_fingerprint = run.source_fingerprint
        created_by_user_id = run.created_by

    else:
        # Sumber 2: Dari hasil tersimpan analysis_results
        req_mode = ModeAnalisa.ai if mode == "ai" else ModeAnalisa.deterministik
        res_row = (
            db.query(AnalysisResult)
            .filter(
                AnalysisResult.document_id == doc.id,
                AnalysisResult.kind == JenisAnalisa.analisa,
                AnalysisResult.mode == req_mode,
            )
            .order_by(AnalysisResult.created_at.desc(), AnalysisResult.id.desc())
            .first()
        )
        if not res_row or not res_row.result:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail={
                    "code": "belum_dianalisa",
                    "message": "Dokumen belum pernah dianalisa pada mode ini. Jalankan analisa terlebih dahulu.",
                },
            )

        raw_result = res_row.result
        actual_mode = res_row.mode.value if hasattr(res_row.mode, "value") else str(res_row.mode)
        ml_version = res_row.ml_version
        created_at = res_row.created_at
        source_fingerprint = res_row.source_fingerprint
        created_by_user_id = res_row.created_by

    # Saring rujukan dinamis
    filtered_result, discarded = saring_rujukan_fn(raw_result, allowed_ids)
    if isinstance(filtered_result, dict):
        filtered_result["disaring_backend"] = discarded

    # Evaluasi status basi
    curr_fingerprint = get_document_fingerprint(doc)
    basi = bool(source_fingerprint and source_fingerprint != curr_fingerprint)

    # Identifikasi username pembuat (Dianalisa Oleh) dan pengekspor (Diekspor Oleh)
    exported_by = None
    if current_user:
        exported_by = getattr(current_user, "username", None)

    dianalisa_oleh = None
    if created_by_user_id is not None:
        if isinstance(created_by_user_id, int):
            u = db.query(User).filter(User.id == created_by_user_id).first()
            if u:
                dianalisa_oleh = u.username
        elif isinstance(created_by_user_id, str):
            dianalisa_oleh = created_by_user_id

    # Format ekspor
    identitas = filtered_result.get("identitas") if isinstance(filtered_result, dict) else None

    if format_type == "docx":
        content_bytes = export_to_docx(
            doc_model=doc,
            result=filtered_result,
            mode=actual_mode,
            ml_version=ml_version,
            created_at=created_at,
            run_id=run_id,
            basi=basi,
            exported_by=exported_by,
            dianalisa_oleh=dianalisa_oleh,
        )
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ext = "docx"

    elif format_type == "csv":
        content_bytes = export_to_csv(result=filtered_result)
        media_type = "text/csv; charset=utf-8"
        ext = "csv"

    elif format_type == "json":
        json_payload = {
            "label": "Draft / Rekomendasi",
            "document_id": doc.id,
            "mode_dipakai": actual_mode,
            "run_id": run_id,
            "basi": basi,
            "dibuat": created_at.isoformat() if created_at else None,
            "diekspor": datetime.now(timezone.utc).isoformat(),
            "hasil": filtered_result,
        }
        content_bytes = json.dumps(json_payload, ensure_ascii=False, indent=2).encode("utf-8")
        media_type = "application/json"
        ext = "json"

    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Format ekspor tidak didukung: '{format_type}'. Format yang didukung: docx, csv, json.",
        )

    filename_ascii, filename_utf8 = generate_export_filename(
        identitas=identitas, document_id=doc.id, mode=actual_mode, ext=ext
    )
    encoded_filename = urllib.parse.quote(filename_utf8)
    headers = {
        "Content-Disposition": f'attachment; filename="{filename_ascii}"; filename*=UTF-8\'\'{encoded_filename}'
    }

    return content_bytes, media_type, headers, basi, actual_mode
