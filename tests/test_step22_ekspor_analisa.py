"""
tests/test_step22_ekspor_analisa.py
Pengujian Ekspor Hasil Analisa Regulasi (Langkah 22, US-40 #62).
Kasus Uji: X01 - X15.
"""
import csv
from datetime import date, datetime, timezone
import io
from typing import Optional
import docx
import httpx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.config import settings
from app.main import app
from app.models.analysis_result import AnalysisResult
from app.models.analysis_run import AnalysisRun
from app.models.audit_log import AuditLog
from app.models.category import Category
from app.models.document import Document
from app.models.enums import (
    JenisAnalisa,
    KlasifikasiAkses,
    ModeAnalisa,
    PeranDokumen,
    StatusKeberlakuan,
    StatusPemrosesan,
)
from app.routers.ml_gateway import get_ml_client
from app.services.analysis_result_service import compute_params_hash, get_document_fingerprint
from app.services.analysis_run_service import compute_result_hash
from app.services.audit_service import EXPORT_ANALYSIS
from app.services.storage_service import StorageService
from tests.conftest import make_pdf


def failing_ml_handler(request: httpx.Request) -> httpx.Response:
    pytest.fail(f"Layanan ML dipanggil secara tidak terduga pada endpoint ekspor: {request.url}")


@pytest.fixture(autouse=True)
def setup_ml_config(monkeypatch: pytest.MonkeyPatch):
    """Pastikan ML transport otomatis melempar kegagalan bila dipanggil."""
    monkeypatch.setattr(settings, "ml_service_url", "http://ml-test:8100")
    monkeypatch.setattr(settings, "ml_cache_enabled", True)
    monkeypatch.setattr(settings, "protect_non_public_when_auth_disabled", True)
    monkeypatch.setattr(settings, "auth_enabled", False)
    app.dependency_overrides[get_ml_client] = lambda: httpx.AsyncClient(transport=httpx.MockTransport(failing_ml_handler))
    yield
    app.dependency_overrides.pop(get_ml_client, None)


def _seed_test_doc(
    db: Session,
    storage: StorageService,
    title: str = "Dokumen Uji Ekspor",
    reg_num: str = "11/POJK.03/2024",
    is_active: bool = True,
    access_classification: KlasifikasiAkses = KlasifikasiAkses.publik,
    processing_status: StatusPemrosesan = StatusPemrosesan.terindeks,
) -> Document:
    cat = db.query(Category).filter(Category.name == "POJK").first()
    pdf_bytes = make_pdf(f"Konten {title}")
    rel_path = storage.save_pdf(pdf_bytes, filename_hint=f"{title.replace(' ', '_').lower()}.pdf", subdir="kb/POJK/2024")

    doc = Document(
        title=title,
        regulation_number=reg_num,
        regulation_type="POJK",
        release_date=date(2024, 6, 15),
        regulation_year=2024,
        file_path_pdf=rel_path,
        file_hash=f"hash_{title.replace(' ', '_').lower()}_{reg_num.replace('/', '_')}",
        file_size_bytes=len(pdf_bytes),
        standardized_filename=f"{title}.pdf",
        access_classification=access_classification,
        document_role=PeranDokumen.corpus_eksisting,
        status_keberlakuan=StatusKeberlakuan.berlaku,
        processing_status=processing_status,
        is_active=is_active,
        category_id=cat.id if cat else None,
        full_text=f"Isi teks lengkap dari {title}",
        extracted_at=datetime(2024, 6, 20, 10, 0, 0, tzinfo=timezone.utc),
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


SAMPLE_RESULT = {
    "versi": "v1.2.0-rules",
    "mode_dipakai": "deterministik",
    "identitas": {
        "jenis": "POJK",
        "nomor": "11",
        "tahun": 2024,
        "tentang": "Penerapan Manajemen Risiko Bank",
        "status": "Berlaku",
        "tanggal_terbit": "2024-06-15",
        "penerbit": "Otoritas Jasa Keuangan",
        "jumlah_pasal": 45,
    },
    "kerangka": [
        {"bab": "I", "judul": "Ketentuan Umum", "pasal": ["1", "2"]},
        {"bab": "II", "judul": "Tata Kelola", "pasal": ["3", "4", "5"]},
    ],
    "ringkasan": [
        {
            "kalimat": "Bank wajib menerapkan manajemen risiko secara efektif dan komprehensif.",
            "rujukan": ["Pasal 2", "Pasal 3"],
        },
        {
            "kalimat": "Komite Pemantau Risiko harus melakukan evaluasi berkala terhadap profil risiko.",
            "rujukan": ["Pasal 10"],
        },
    ],
    "poin_utama": [
        {
            "id": 1,
            "kategori": "Kewajiban",
            "pasal": "2",
            "halaman": 5,
            "teks": "Penyusunan kerangka manajemen risiko wajib disetujui oleh Direksi dan Dewan Komisaris.",
            "terpotong": False,
        },
        {
            "id": 2,
            "kategori": "Larangan",
            "pasal": "7",
            "halaman": 12,
            "teks": "Bank dilarang mengabaikan limit risiko kredit yang telah ditetapkan dalam rencana bisnis.",
            "terpotong": True,
        },
        {
            "id": 3,
            "kategori": "Pelaporan",
            "pasal": "15",
            "halaman": 20,
            "teks": "Laporan profil risiko triwulanan wajib disampaikan kepada OJK paling lambat tanggal 15.",
            "terpotong": False,
        },
    ],
    "dasar_hukum": [
        "Undang-Undang Nomor 4 Tahun 2023 tentang Pengembangan dan Penguatan Sektor Keuangan",
    ],
    "sanksi": [
        "Sanksi administratif berupa teguran tertulis dan/atau denda",
    ],
    "catatan": "Hasil analisa otomatis menggunakan aturan hukum perbankan.",
}


SAMPLE_AI_RESULT = {
    **SAMPLE_RESULT,
    "mode_dipakai": "ai",
    "ringkasan_ai": "Narasi ringkasan komprehensif berbantuan kecerdasan buatan untuk manajemen risiko.",
    "poin_utama": [
        {
            "id": 1,
            "kategori": "Kewajiban",
            "pasal": "2",
            "halaman": 5,
            "teks": "Penyusunan kerangka manajemen risiko wajib disetujui oleh Direksi dan Dewan Komisaris.",
            "terpotong": False,
            "parafrase": "Direksi serta Komisaris bank bertanggung jawab penuh mengesahkan arsitektur manajemen risiko.",
        },
        {
            "id": 2,
            "kategori": "Larangan",
            "pasal": "7",
            "halaman": 12,
            "teks": "Bank dilarang mengabaikan limit risiko kredit yang telah ditetapkan dalam rencana bisnis.",
            "terpotong": True,
            "parafrase": "Pelanggaran batas toleransi risiko kredit tidak diizinkan.",
        },
        {
            "id": 3,
            "kategori": "Pelaporan",
            "pasal": "15",
            "halaman": 20,
            "teks": "Laporan profil risiko triwulanan wajib disampaikan kepada OJK paling lambat tanggal 15.",
            "terpotong": False,
            "parafrase": "Penyampaian profil risiko kepada OJK wajib tepat waktu setiap triwulan.",
        },
    ],
}


def _seed_analysis_result(
    db: Session,
    doc: Document,
    mode: ModeAnalisa = ModeAnalisa.deterministik,
    raw_result: Optional[dict] = None,
    created_at: Optional[datetime] = None,
) -> AnalysisResult:
    res_dict = raw_result if raw_result is not None else SAMPLE_RESULT
    row = AnalysisResult(
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode=mode,
        params_hash=compute_params_hash(JenisAnalisa.analisa),
        source_fingerprint=get_document_fingerprint(doc),
        result=res_dict,
        ml_version=res_dict.get("versi", "1.0"),
        duration_ms=1200,
        created_at=created_at or datetime.now(timezone.utc),
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def test_x01_export_docx_structure(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X01: format=docx -> 200, docx memuat label DRAFT / REKOMENDASI, kalimat ringkasan, teks poin, dan Pasal rujukan."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X01", reg_num="01/X01/2024")
    _seed_analysis_result(db_session, doc)

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=docx")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in resp.headers["content-type"]

    doc_file = docx.Document(io.BytesIO(resp.content))
    full_text = "\n".join(
        [p.text for p in doc_file.paragraphs]
        + [cell.text for t in doc_file.tables for row in t.rows for cell in row.cells]
    )

    # 1. Label
    assert "DRAFT / REKOMENDASI" in full_text
    # 2. Kalimat ringkasan
    assert "Bank wajib menerapkan manajemen risiko secara efektif dan komprehensif." in full_text
    assert "Komite Pemantau Risiko harus melakukan evaluasi berkala terhadap profil risiko." in full_text
    # 3. Poin utama
    assert "Penyusunan kerangka manajemen risiko wajib disetujui oleh Direksi dan Dewan Komisaris." in full_text
    assert "Bank dilarang mengabaikan limit risiko kredit yang telah ditetapkan dalam rencana bisnis. (dipotong)" in full_text
    assert "Laporan profil risiko triwulanan wajib disampaikan kepada OJK paling lambat tanggal 15." in full_text
    # 4. Rujukan pasal
    assert "Pasal 2" in full_text
    assert "Pasal 7" in full_text
    assert "Pasal 10" in full_text
    assert "Pasal 15" in full_text


def test_x02_export_csv_structure(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X02: format=csv -> 200, BOM, baris label, 2 ringkasan dg pasal_rujukan, 3 poin_utama dg pasal_rujukan & halaman."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X02", reg_num="02/X02/2024")
    _seed_analysis_result(db_session, doc)

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=csv")
    assert resp.status_code == 200
    assert resp.content.startswith(b"\xef\xbb\xbf")

    csv_text = resp.content.decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(csv_text)))

    # Baris label
    label_rows = [r for r in rows if r["bagian"] == "label"]
    assert len(label_rows) == 1
    assert "DRAFT / REKOMENDASI" in label_rows[0]["teks"]

    # 2 baris ringkasan dengan pasal_rujukan terisi
    ringkasan_rows = [r for r in rows if r["bagian"] == "ringkasan"]
    assert len(ringkasan_rows) == 2
    for r in ringkasan_rows:
        assert len(r["pasal_rujukan"].strip()) > 0

    # 3 baris poin_utama dengan pasal_rujukan dan halaman terisi
    poin_rows = [r for r in rows if r["bagian"] == "poin_utama"]
    assert len(poin_rows) == 3
    for r in poin_rows:
        assert len(r["pasal_rujukan"].strip()) > 0
        assert len(r["halaman"].strip()) > 0


def test_x03_export_json_structure(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X03: format=json -> 200, label, hasil.poin_utama panjang 3, basi=false."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X03", reg_num="03/X03/2024")
    _seed_analysis_result(db_session, doc)

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=json")
    assert resp.status_code == 200
    data = resp.json()

    assert data["label"] == "Draft / Rekomendasi"
    assert data["document_id"] == doc.id
    assert data["mode_dipakai"] == "deterministik"
    assert data["basi"] is False
    assert len(data["hasil"]["poin_utama"]) == 3


def test_x04_not_yet_analyzed_returns_404(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X04: Dokumen belum pernah dianalisa -> 404 belum_dianalisa."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X04", reg_num="04/X04/2024")

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor")
    assert resp.status_code == 404
    data = resp.json()
    assert data["detail"]["code"] == "belum_dianalisa"


def test_x05_mode_ai_requested_but_only_deterministik_exists(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X05: mode=ai tetapi hanya ada hasil deterministik -> 404 belum_dianalisa."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X05", reg_num="05/X05/2024")
    _seed_analysis_result(db_session, doc, mode=ModeAnalisa.deterministik)

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?mode=ai")
    assert resp.status_code == 404
    data = resp.json()
    assert data["detail"]["code"] == "belum_dianalisa"


def test_x06_export_ai_with_ringkasan_ai_and_parafrase(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X06: Hasil AI dengan ringkasan_ai + parafrase -> docx memuat ringkasan deterministik & subbagian AI."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X06", reg_num="06/X06/2024")
    _seed_analysis_result(db_session, doc, mode=ModeAnalisa.ai, raw_result=SAMPLE_AI_RESULT)

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?mode=ai&format=docx")
    assert resp.status_code == 200

    doc_file = docx.Document(io.BytesIO(resp.content))
    full_text = "\n".join([p.text for p in doc_file.paragraphs])

    # Ringkasan deterministik tetap ada
    assert "Bank wajib menerapkan manajemen risiko secara efektif dan komprehensif." in full_text
    # Subbagian Narasi AI-Assisted ada
    assert "Narasi AI-Assisted" in full_text
    assert "Narasi ringkasan komprehensif berbantuan kecerdasan buatan" in full_text
    # Parafrase AI ada
    assert "Parafrase AI:" in full_text
    assert "Direksi serta Komisaris bank bertanggung jawab penuh" in full_text


def test_x07_run_id_from_cache_run_resolves_result_hash(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X07: run_id dari baris sumber=cache -> hasil diresolusi lewat result_hash; isi sama dengan run ml."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X07", reg_num="07/X07/2024")
    res_hash = compute_result_hash(SAMPLE_RESULT)

    # 1. Run ML
    run_ml = AnalysisRun(
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode_diminta=ModeAnalisa.deterministik,
        mode_dipakai=ModeAnalisa.deterministik,
        sumber="ml",
        segar=False,
        params={"mode": "deterministik"},
        params_hash="hash_x07",
        source_fingerprint=get_document_fingerprint(doc),
        ml_version="1.0",
        result_hash=res_hash,
        result=SAMPLE_RESULT,
        http_status=200,
    )
    db_session.add(run_ml)
    db_session.commit()

    # 2. Run Cache (result = NULL)
    run_cache = AnalysisRun(
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode_diminta=ModeAnalisa.deterministik,
        mode_dipakai=ModeAnalisa.deterministik,
        sumber="cache",
        segar=False,
        params={"mode": "deterministik"},
        params_hash="hash_x07",
        source_fingerprint=get_document_fingerprint(doc),
        ml_version="1.0",
        result_hash=res_hash,
        result=None,
        http_status=200,
    )
    db_session.add(run_cache)
    db_session.commit()

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?run_id={run_cache.id}&format=json")
    assert resp.status_code == 200
    data = resp.json()
    assert data["run_id"] == run_cache.id
    assert data["hasil"]["ringkasan"] == SAMPLE_RESULT["ringkasan"]


def test_x08_run_id_different_document_or_failed_returns_404(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X08: run_id milik dokumen lain -> 404 / sumber=gagal -> 404 hasil_tidak_tersedia."""
    doc_a = _seed_test_doc(db_session, test_storage, title="Doc A X08", reg_num="08/A/2024")
    doc_b = _seed_test_doc(db_session, test_storage, title="Doc B X08", reg_num="08/B/2024")

    run_b = AnalysisRun(
        document_id=doc_b.id,
        kind=JenisAnalisa.analisa,
        mode_diminta=ModeAnalisa.deterministik,
        mode_dipakai=ModeAnalisa.deterministik,
        sumber="ml",
        segar=False,
        params={"mode": "deterministik"},
        params_hash="hash_b",
        source_fingerprint=get_document_fingerprint(doc_b),
        result=SAMPLE_RESULT,
        http_status=200,
    )
    # Run gagal di doc_a
    run_failed_a = AnalysisRun(
        document_id=doc_a.id,
        kind=JenisAnalisa.analisa,
        mode_diminta=ModeAnalisa.deterministik,
        mode_dipakai=None,
        sumber="gagal",
        segar=False,
        params={"mode": "deterministik"},
        params_hash="hash_a",
        source_fingerprint=get_document_fingerprint(doc_a),
        result=None,
        http_status=503,
        error_code="layanan_analisa_tidak_tersedia",
    )
    db_session.add_all([run_b, run_failed_a])
    db_session.commit()

    # 1. Akses run_b lewat doc_a -> 404
    r1 = client.get(f"/api/v1/documents/{doc_a.id}/analisa/ekspor?run_id={run_b.id}")
    assert r1.status_code == 404

    # 2. Akses run_failed_a -> 404 hasil_tidak_tersedia
    r2 = client.get(f"/api/v1/documents/{doc_a.id}/analisa/ekspor?run_id={run_failed_a.id}")
    assert r2.status_code == 404
    assert r2.json()["detail"]["code"] == "hasil_tidak_tersedia"


def test_x09_extracted_at_changed_shows_stale_warning(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X09: extracted_at dokumen berubah setelah analisa -> docx memuat peringatan; json basi=true."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X09", reg_num="09/X09/2024")
    _seed_analysis_result(db_session, doc)

    # Ubah extracted_at dokumen
    doc.extracted_at = datetime(2026, 10, 18, 12, 0, 0, tzinfo=timezone.utc)
    db_session.commit()

    # 1. docx
    resp_docx = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=docx")
    assert resp_docx.status_code == 200
    doc_file = docx.Document(io.BytesIO(resp_docx.content))
    full_text = "\n".join([p.text for p in doc_file.paragraphs])
    assert "Dokumen telah berubah sejak analisa ini dibuat — jalankan analisa ulang." in full_text

    # 2. json
    resp_json = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=json")
    assert resp_json.status_code == 200
    assert resp_json.json()["basi"] is True


def test_x10_filtered_deactivated_reference(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X10: Hasil memuat rujukan ke dokumen yang kemudian dinonaktifkan -> rujukan tidak muncul, disaring_backend >= 1."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X10", reg_num="10/X10/2024")
    ref_doc = _seed_test_doc(db_session, test_storage, title="Ref Doc X10", reg_num="10/REF/2024")

    res_with_ref = {
        **SAMPLE_RESULT,
        "rujukan": [
            {"document_id": ref_doc.id, "pasal": "12", "pdf": "ref.pdf", "kutipan": "Kutipan rahasia"},
        ],
    }
    _seed_analysis_result(db_session, doc, raw_result=res_with_ref)

    # Nonaktifkan dokumen rujukan
    ref_doc.is_active = False
    db_session.commit()

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=json")
    assert resp.status_code == 200
    data = resp.json()
    assert data["hasil"]["disaring_backend"] >= 1
    # document_id rujukan dinetralkan
    assert data["hasil"]["rujukan"][0]["document_id"] is None
    assert "kutipan" not in data["hasil"]["rujukan"][0]


def test_x11_non_public_protected_mode_returns_403(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X11: Dokumen non-publik, mode lindung -> 403."""
    doc_np = _seed_test_doc(
        db_session,
        test_storage,
        title="Doc NP X11",
        reg_num="11/NP/2024",
        access_classification=KlasifikasiAkses.non_publik,
    )

    resp = client.get(f"/api/v1/documents/{doc_np.id}/analisa/ekspor")
    assert resp.status_code == 403


def test_x12_content_disposition_header(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X12: Header Content-Disposition -> memuat filename*= dan analisa_POJK_11_2024_deterministik.docx."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X12", reg_num="12/X12/2024")
    _seed_analysis_result(db_session, doc)

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=docx")
    assert resp.status_code == 200
    cd = resp.headers.get("content-disposition", "")
    assert "filename*=" in cd
    assert "analisa_POJK_11_2024_deterministik.docx" in cd


def test_x13_audit_entry_created(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X13: Audit -> satu entri EXPORT_ANALYSIS dengan detail.format dan detail.mode."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X13", reg_num="13/X13/2024")
    _seed_analysis_result(db_session, doc)

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=csv")
    assert resp.status_code == 200

    audit = (
        db_session.query(AuditLog)
        .filter(AuditLog.action == EXPORT_ANALYSIS, AuditLog.target_resource == f"document:{doc.id}")
        .first()
    )
    assert audit is not None
    assert audit.detail.get("format") == "csv"
    assert audit.detail.get("mode") == "deterministik"


def test_x14_result_with_missing_fields_renders_gracefully(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X14: Hasil dengan field hilang (tanpa identitas, dasar_hukum, sanksi) -> 200 untuk docx, csv, json tanpa error."""
    doc = _seed_test_doc(db_session, test_storage, title="Doc X14", reg_num="14/X14/2024")
    minimal_result = {
        "versi": "1.0",
        "mode_dipakai": "deterministik",
        "ringkasan": [],
        "poin_utama": [],
    }
    _seed_analysis_result(db_session, doc, raw_result=minimal_result)

    # 1. docx
    r_docx = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=docx")
    assert r_docx.status_code == 200

    # 2. csv
    r_csv = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=csv")
    assert r_csv.status_code == 200

    # 3. json
    r_json = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=json")
    assert r_json.status_code == 200


def test_x15_ml_never_called(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """X15: ML -> handler MockTransport failing_ml_handler tidak pernah dipanggil di seluruh pengujian."""
    # Dipastikan oleh autouse fixture setup_ml_config dengan failing_ml_handler
    doc = _seed_test_doc(db_session, test_storage, title="Doc X15", reg_num="15/X15/2024")
    _seed_analysis_result(db_session, doc)

    # Lakukan ekspor ketiga format
    client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=docx")
    client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=csv")
    client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=json")
