"""
tests/test_step25_cache_usang.py
Pengujian Invalidasi Cache Analisa setelah Pembaruan Mesin ML (Langkah 25).
Kasus Uji: V01 - V13.
"""
from datetime import date, datetime, timezone
import io
import docx
import httpx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.config import settings
from app.main import app
from app.models.analysis_result import AnalysisResult
from app.models.analysis_run import AnalysisRun
from app.models.audit_log import AuditLog
from app.models.category import Category
from app.models.document import Document
from app.models.enums import (
    JenisAnalisa,
    KlasifikasiAkses,
    ModeAnalisa,
    PeranDokumen,
    StatusKeberlakuan,
    StatusPemrosesan,
)
from app.routers.ml_gateway import get_ml_client
from app.services.analysis_result_service import compute_params_hash, get_document_fingerprint, save_analysis_result
from app.services.audit_service import ANALYSIS_CACHE_INVALIDATED
from app.services.storage_service import StorageService
from tests.conftest import make_pdf


@pytest.fixture(autouse=True)
def setup_ml_config(monkeypatch: pytest.MonkeyPatch):
    """Pastikan konfigurasi ML aktif dan terisolasi untuk tiap pengujian."""
    monkeypatch.setattr(settings, "ml_service_url", "http://ml-test:8100")
    monkeypatch.setattr(settings, "ml_cache_enabled", True)
    monkeypatch.setattr(settings, "protect_non_public_when_auth_disabled", True)
    monkeypatch.setattr(settings, "auth_enabled", False)
    yield
    app.dependency_overrides.pop(get_ml_client, None)


def _seed_test_doc(
    db: Session,
    storage: StorageService,
    title: str = "Dokumen Uji Cache Usang",
    reg_num: str = "25/POJK.03/2026",
    is_active: bool = True,
    access_classification: KlasifikasiAkses = KlasifikasiAkses.publik,
    processing_status: StatusPemrosesan = StatusPemrosesan.terindeks,
) -> Document:
    cat = db.query(Category).filter(Category.name == "POJK").first()
    pdf_bytes = make_pdf(f"Konten {title}")
    rel_path = storage.save_pdf(pdf_bytes, filename_hint=f"{title.replace(' ', '_').lower()}.pdf", subdir="kb/POJK/2026")

    doc = Document(
        title=title,
        regulation_number=reg_num,
        regulation_type="POJK",
        release_date=date(2026, 1, 1),
        regulation_year=2026,
        file_path_pdf=rel_path,
        file_hash=f"hash_{title.replace(' ', '_').lower()}",
        file_size_bytes=len(pdf_bytes),
        standardized_filename=f"{title}.pdf",
        access_classification=access_classification,
        document_role=PeranDokumen.corpus_eksisting,
        status_keberlakuan=StatusKeberlakuan.berlaku,
        processing_status=processing_status,
        is_active=is_active,
        category_id=cat.id if cat else None,
        full_text=f"Isi teks lengkap dari {title}",
        extracted_at=datetime(2026, 10, 1, 10, 0, 0, tzinfo=timezone.utc),
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


def _headers() -> dict:
    return {"X-Internal-API-Key": settings.internal_api_key}


def test_v01_invalidate_forces_recomputation(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V01: Analisa dua kali, lalu invalidate (tanpa filter), lalu analisa lagi:
         panggilan ML ke-1 dan ke-3 terjadi (total 2), panggilan ke-2 cache hit; respons ke-3 cache.hit=false.
    """
    doc = _seed_test_doc(db_session, test_storage, title="Dokumen V01", reg_num="01/V01/2026")
    call_count = [0]

    def handler(request: httpx.Request) -> httpx.Response:
        call_count[0] += 1
        ver = "v2.0" if call_count[0] == 1 else "v2.1"
        return httpx.Response(
            200,
            json={
                "versi": ver,
                "mode_dipakai": "deterministik",
                "ringkasan": f"Ringkasan {ver}",
                "poin_kunci": ["Poin A"],
            },
        )

    app.dependency_overrides[get_ml_client] = lambda: httpx.AsyncClient(transport=httpx.MockTransport(handler))

    # Panggilan ke-1 (ML call 1)
    resp1 = client.get(f"/api/v1/ml/analisa/{doc.id}")
    assert resp1.status_code == 200
    assert call_count[0] == 1
    assert resp1.json()["cache"]["hit"] is False

    # Panggilan ke-2 (cache hit)
    resp2 = client.get(f"/api/v1/ml/analisa/{doc.id}")
    assert resp2.status_code == 200
    assert call_count[0] == 1
    assert resp2.json()["cache"]["hit"] is True

    # Invalidate tanpa filter
    inv_resp = client.post("/api/v1/internal/analysis-cache/invalidate", json={}, headers=_headers())
    assert inv_resp.status_code == 200
    assert inv_resp.json()["ditandai"] == 1

    # Panggilan ke-3 (ML call 2) -> karena usang, cache miss
    resp3 = client.get(f"/api/v1/ml/analisa/{doc.id}")
    assert resp3.status_code == 200
    assert call_count[0] == 2
    assert resp3.json()["cache"]["hit"] is False
    assert resp3.json()["ringkasan"] == "Ringkasan v2.1"


def test_v02_save_analysis_result_clears_usang_sejak_and_updates_ml_version(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V02: Setelah V01 (atau simpan hasil analisa sukses), baris usang_sejak kembali NULL dan ml_version diperbarui.
    """
    doc = _seed_test_doc(db_session, test_storage, title="Dokumen V02", reg_num="02/V02/2026")
    params_hash = compute_params_hash(JenisAnalisa.analisa)
    fp = get_document_fingerprint(doc)

    # 1. Simpan hasil lama
    row = save_analysis_result(
        db=db_session,
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=fp,
        raw_result={"versi": "v2.0"},
        ml_version="v2.0",
    )
    # Tandai usang
    row.usang_sejak = datetime.now(timezone.utc)
    row.alasan_usang = "mesin diperbarui"
    db_session.commit()
    db_session.refresh(row)
    assert row.usang_sejak is not None
    assert row.alasan_usang == "mesin diperbarui"

    # 2. Simpan hasil analisa baru (upsert sukses)
    updated_row = save_analysis_result(
        db=db_session,
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=fp,
        raw_result={"versi": "v2.1", "ringkasan": "Hasil Baru v2.1"},
        ml_version="v2.1",
    )

    # 3. Verifikasi usang_sejak & alasan_usang kembali NULL dan ml_version ter-update
    assert updated_row.usang_sejak is None
    assert updated_row.alasan_usang is None
    assert updated_row.ml_version == "v2.1"


def test_v03_filter_versi_mesin(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V03: versi_mesin="v2.1"; dok A tersimpan v2.0, dok B tersimpan v2.1 -> ditandai=1; hanya A yang usang.
    """
    doc_a = _seed_test_doc(db_session, test_storage, title="Dokumen V03 A", reg_num="03A/V03/2026")
    doc_b = _seed_test_doc(db_session, test_storage, title="Dokumen V03 B", reg_num="03B/V03/2026")

    fp_a = get_document_fingerprint(doc_a)
    fp_b = get_document_fingerprint(doc_b)
    params_hash = compute_params_hash(JenisAnalisa.analisa)

    save_analysis_result(
        db=db_session,
        document_id=doc_a.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=fp_a,
        raw_result={"versi": "v2.0", "mode_dipakai": "deterministik"},
        ml_version="v2.0",
    )
    save_analysis_result(
        db=db_session,
        document_id=doc_b.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=fp_b,
        raw_result={"versi": "v2.1", "mode_dipakai": "deterministik"},
        ml_version="v2.1",
    )

    resp = client.post(
        "/api/v1/internal/analysis-cache/invalidate",
        json={"versi_mesin": "v2.1"},
        headers=_headers(),
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["ditandai"] == 1
    assert data["versi_mesin"] == "v2.1"

    row_a = db_session.query(AnalysisResult).filter(AnalysisResult.document_id == doc_a.id).first()
    row_b = db_session.query(AnalysisResult).filter(AnalysisResult.document_id == doc_b.id).first()
    assert row_a.usang_sejak is not None
    assert row_b.usang_sejak is None


def test_v04_invalidate_twice_idempotency(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V04: Invalidate dua kali berturut-turut -> panggilan ke-2: ditandai=0, sudah_usang_sebelumnya >= 1.
    """
    doc = _seed_test_doc(db_session, test_storage, title="Dokumen V04", reg_num="04/V04/2026")
    fp = get_document_fingerprint(doc)
    params_hash = compute_params_hash(JenisAnalisa.analisa)
    save_analysis_result(
        db=db_session,
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=fp,
        raw_result={"versi": "v2.0"},
        ml_version="v2.0",
    )

    # Panggilan ke-1
    resp1 = client.post("/api/v1/internal/analysis-cache/invalidate", json={}, headers=_headers())
    assert resp1.status_code == 200
    assert resp1.json()["ditandai"] >= 1

    # Panggilan ke-2
    resp2 = client.post("/api/v1/internal/analysis-cache/invalidate", json={}, headers=_headers())
    assert resp2.status_code == 200
    data2 = resp2.json()
    assert data2["ditandai"] == 0
    assert data2["sudah_usang_sebelumnya"] >= 1


def test_v05_filter_document_ids(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V05: document_ids=[A] -> hanya baris A yang ditandai.
    """
    doc_a = _seed_test_doc(db_session, test_storage, title="Dokumen V05 A", reg_num="05A/V05/2026")
    doc_b = _seed_test_doc(db_session, test_storage, title="Dokumen V05 B", reg_num="05B/V05/2026")
    params_hash = compute_params_hash(JenisAnalisa.analisa)

    save_analysis_result(
        db=db_session,
        document_id=doc_a.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=get_document_fingerprint(doc_a),
        raw_result={"versi": "v2.0"},
        ml_version="v2.0",
    )
    save_analysis_result(
        db=db_session,
        document_id=doc_b.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=get_document_fingerprint(doc_b),
        raw_result={"versi": "v2.0"},
        ml_version="v2.0",
    )

    resp = client.post(
        "/api/v1/internal/analysis-cache/invalidate",
        json={"document_ids": [doc_a.id]},
        headers=_headers(),
    )
    assert resp.status_code == 200
    assert resp.json()["ditandai"] == 1

    row_a = db_session.query(AnalysisResult).filter(AnalysisResult.document_id == doc_a.id).first()
    row_b = db_session.query(AnalysisResult).filter(AnalysisResult.document_id == doc_b.id).first()
    assert row_a.usang_sejak is not None
    assert row_b.usang_sejak is None


def test_v06_filter_kind_harmonisasi(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V06: kind="harmonisasi" -> baris analisa tidak tersentuh.
    """
    doc = _seed_test_doc(db_session, test_storage, title="Dokumen V06", reg_num="06/V06/2026")
    params_hash = compute_params_hash(JenisAnalisa.analisa)
    save_analysis_result(
        db=db_session,
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=get_document_fingerprint(doc),
        raw_result={"versi": "v2.0"},
        ml_version="v2.0",
    )

    resp = client.post(
        "/api/v1/internal/analysis-cache/invalidate",
        json={"kind": "harmonisasi"},
        headers=_headers(),
    )
    assert resp.status_code == 200
    assert resp.json()["ditandai"] == 0

    row = db_session.query(AnalysisResult).filter(AnalysisResult.document_id == doc.id).first()
    assert row.usang_sejak is None


def test_v07_empty_document_ids_returns_422(client: TestClient):
    """
    V07: document_ids=[] -> 422 Unprocessable Entity.
    """
    resp = client.post(
        "/api/v1/internal/analysis-cache/invalidate",
        json={"document_ids": []},
        headers=_headers(),
    )
    assert resp.status_code == 422


def test_v08_missing_or_invalid_api_key_returns_401(client: TestClient):
    """
    V08: Tanpa / salah X-Internal-API-Key -> 401 Unauthorized.
    """
    resp1 = client.post("/api/v1/internal/analysis-cache/invalidate", json={})
    assert resp1.status_code == 401

    resp2 = client.post(
        "/api/v1/internal/analysis-cache/invalidate",
        json={},
        headers={"X-Internal-API-Key": "kunci-palsu-salah"},
    )
    assert resp2.status_code == 401


def test_v09_stale_cache_fallback_when_ml_offline(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V09: Baris usang, lalu ML mati (ConnectError) -> 200, hasil lama tersaji, cache.basi=true.
    """
    doc = _seed_test_doc(db_session, test_storage, title="Dokumen V09", reg_num="09/V09/2026")
    params_hash = compute_params_hash(JenisAnalisa.analisa)
    row = save_analysis_result(
        db=db_session,
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=get_document_fingerprint(doc),
        raw_result={"versi": "v2.0", "mode_dipakai": "deterministik", "ringkasan": "Hasil lama sebelum ML mati"},
        ml_version="v2.0",
    )
    row.usang_sejak = datetime.now(timezone.utc)
    row.alasan_usang = "mesin v2.1"
    db_session.commit()

    def handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("Layanan ML tidak dapat dihubungi")

    app.dependency_overrides[get_ml_client] = lambda: httpx.AsyncClient(transport=httpx.MockTransport(handler))

    resp = client.get(f"/api/v1/ml/analisa/{doc.id}")
    assert resp.status_code == 200
    data = resp.json()
    assert data["cache"]["hit"] is True
    assert data["cache"]["basi"] is True
    assert data["ringkasan"] == "Hasil lama sebelum ML mati"


def test_v10_analisa_banding_shows_basi_true_for_stale_cache(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V10: /analisa/banding dengan baris usang -> basi=true.
    """
    doc = _seed_test_doc(db_session, test_storage, title="Dokumen V10", reg_num="10/V10/2026")
    params_hash = compute_params_hash(JenisAnalisa.analisa)
    row = save_analysis_result(
        db=db_session,
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=get_document_fingerprint(doc),
        raw_result={"versi": "v2.0", "ringkasan": "Banding V10"},
        ml_version="v2.0",
    )
    row.usang_sejak = datetime.now(timezone.utc)
    db_session.commit()

    resp = client.get(f"/api/v1/documents/{doc.id}/analisa/banding")
    assert resp.status_code == 200
    data = resp.json()
    assert data["deterministik"] is not None
    assert data["deterministik"]["basi"] is True


def test_v11_export_docx_and_json_contains_stale_warning(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V11: Ekspor DOCX baris usang memuat kalimat 'Mesin analisa telah diperbarui…', JSON memuat basi=true & pesan usang.
    """
    doc = _seed_test_doc(db_session, test_storage, title="Dokumen V11", reg_num="11/V11/2026")
    params_hash = compute_params_hash(JenisAnalisa.analisa)
    row = save_analysis_result(
        db=db_session,
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=get_document_fingerprint(doc),
        raw_result={
            "identitas": {"jenis": "POJK", "nomor": "11", "tahun": 2026, "tentang": "Dokumen V11"},
            "ringkasan": [{"kalimat": "Ringkasan V11", "rujukan": []}],
            "poin_utama": [{"kategori": "Kewajiban", "teks": "Wajib lapor", "pasal": "1"}],
        },
        ml_version="v2.0",
    )
    row.usang_sejak = datetime.now(timezone.utc)
    row.alasan_usang = "mesin v2.1"
    db_session.commit()

    # 1. Ekspor DOCX
    resp_docx = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=docx")
    assert resp_docx.status_code == 200
    doc_file = docx.Document(io.BytesIO(resp_docx.content))
    full_text = "\n".join([p.text for p in doc_file.paragraphs])
    assert "Mesin analisa telah diperbarui sejak hasil ini dibuat — jalankan analisa ulang." in full_text

    # 2. Ekspor JSON
    resp_json = client.get(f"/api/v1/documents/{doc.id}/analisa/ekspor?format=json")
    assert resp_json.status_code == 200
    data_json = resp_json.json()
    assert data_json["basi"] is True
    assert "Mesin analisa telah diperbarui sejak hasil ini dibuat — jalankan analisa ulang." in resp_json.text


def test_v12_analysis_runs_count_unchanged_after_invalidation(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V12: Jumlah baris analysis_runs sebelum dan sesudah invalidate tidak berubah.
    """
    doc = _seed_test_doc(db_session, test_storage, title="Dokumen V12", reg_num="12/V12/2026")
    params_hash = compute_params_hash(JenisAnalisa.analisa)
    save_analysis_result(
        db=db_session,
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=get_document_fingerprint(doc),
        raw_result={"versi": "v2.0"},
        ml_version="v2.0",
    )

    runs_before = db_session.query(AnalysisRun).count()

    resp = client.post("/api/v1/internal/analysis-cache/invalidate", json={}, headers=_headers())
    assert resp.status_code == 200

    runs_after = db_session.query(AnalysisRun).count()
    assert runs_after == runs_before


def test_v13_audit_log_analysis_cache_invalidated(
    client: TestClient, db_session: Session, test_storage: StorageService
):
    """
    V13: Audit -> satu entri ANALYSIS_CACHE_INVALIDATED dengan detail.ditandai, target_resource="analysis_results", user_id=None.
    """
    doc = _seed_test_doc(db_session, test_storage, title="Dokumen V13", reg_num="13/V13/2026")
    params_hash = compute_params_hash(JenisAnalisa.analisa)
    save_analysis_result(
        db=db_session,
        document_id=doc.id,
        kind=JenisAnalisa.analisa,
        mode=ModeAnalisa.deterministik,
        params_hash=params_hash,
        source_fingerprint=get_document_fingerprint(doc),
        raw_result={"versi": "v2.0"},
        ml_version="v2.0",
    )

    resp = client.post(
        "/api/v1/internal/analysis-cache/invalidate",
        json={"versi_mesin": "v2.1", "alasan": "mesin v2.1: uji audit"},
        headers=_headers(),
    )
    assert resp.status_code == 200
    ditandai_count = resp.json()["ditandai"]

    audit = (
        db_session.query(AuditLog)
        .filter(AuditLog.action == ANALYSIS_CACHE_INVALIDATED)
        .order_by(AuditLog.id.desc())
        .first()
    )
    assert audit is not None
    assert audit.target_resource == "analysis_results"
    assert audit.user_id is None
    assert audit.detail is not None
    assert audit.detail.get("ditandai") == ditandai_count
    assert audit.detail.get("versi_mesin") == "v2.1"
    assert audit.detail.get("alasan") == "mesin v2.1: uji audit"
